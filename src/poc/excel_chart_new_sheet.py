from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

# Create workbook and main sheet
wb = Workbook()
data_ws = wb.active
data_ws.title = "Data"

# Add data
data = [
    ["X", "Y"],
    [1, 10],
    [2, 20],
    [3, 30],
    [4, 40],
]
for row in data:
    data_ws.append(row)

# Create chart
chart = LineChart()
chart.title = "Simple Line Chart"

# Y values (with header for title)
data_ref = Reference(data_ws, min_col=2, min_row=1, max_row=5)
# X values (no header)
cats = Reference(data_ws, min_col=1, min_row=2, max_row=5)

# Add data to chart
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats)

# Create a new sheet for the chart
chart_ws = wb.create_sheet(title="Chart")
chart_ws.add_chart(chart, "A1")

# Save the workbook
wb.save("/home/pete/Documents/testdata/minimal_line_chart.xlsx")
# -*- coding: utf-8 -*-

