import argparse
import csv
import os
import re
import logging
from datetime import datetime
from pprint import pprint
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import NoSuchElementException, TimeoutException, WebDriverException

from file_writer import write as f_write
from excel_interface import xlsx_report

import pandas as pd
import matplotlib.pyplot as plt
import os

# Logging configuration
logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s')


def setup_driver():
    options = webdriver.FirefoxOptions()
    options.add_argument("--width=1200")
    options.add_argument("--height=800")
    options.add_argument("--headless")
    return webdriver.Firefox(options=options)

def login_to_zwiftpower(driver, email, password):
    try:
        driver.get("https://zwiftpower.com/")
        WebDriverWait(driver, 10).until(EC.element_to_be_clickable((By.LINK_TEXT, "Login with Zwift"))).click()

        WebDriverWait(driver, 20).until(EC.presence_of_element_located((By.NAME, "username")))
        driver.find_element(By.NAME, "username").send_keys(email)
        driver.find_element(By.NAME, "password").send_keys(password)
        driver.find_element(By.XPATH, "//button[@type='submit']").click()

        WebDriverWait(driver, 30).until(EC.url_contains("zwiftpower.com"))
        logging.info("✅ Login successful.")

    except (TimeoutException, WebDriverException, NoSuchElementException) as e:
        logging.error(f"❌ Login failed: {e}")
        raise

def scrape_profile_data(driver, profile_url):
    driver.get(profile_url)

    WebDriverWait(driver, 20).until(
        EC.presence_of_element_located((By.XPATH, "//th[contains(text(), 'zFTP')]"))
    )

    data = {
        "date": datetime.now().strftime("%Y-%m-%d"),
        "weight": None,
        "zftp": None,
        "zftp_wkg": None,
        "racing_score": None,
        "15s_w": None, "1m_w": None, "5m_w": None, "20m_w": None,
        "15s_wkg": None, "1m_wkg": None, "5m_wkg": None, "20m_wkg": None,
    }

    try:
        rows = driver.find_elements(By.XPATH, "//tr[th and td]")
        for row in rows:
            label = row.find_element(By.XPATH, "./th").text.strip().lower()
            value = row.find_element(By.XPATH, "./td").text.strip()

            if "zftp" in label:
                data["zftp"] = value
            elif "weight" in label:
                data["weight"] = value
                data["weight"] = re.sub(r"[^\d.]", "", data["weight"])
            elif "racing score" in label:
                data["racing_score"] = re.search(r"\d+", value).group()


        # Strip "w" from zftp
        if data["zftp"] and data["zftp"].lower().endswith("w"):
            data["zftp"] = re.sub(r"[^\d.]", "", data["zftp"])

        # Calculate zftp_wkg
        try:
            zftp_float = float(data["zftp"])
            weight_float = float(data["weight"])
            data["zftp_wkg"] = round(zftp_float / weight_float, 2)
        except (ValueError, TypeError, ZeroDivisionError):
            data["zftp_wkg"] = None

        source = driver.page_source
        watts = extract_power(source, "w")
        wkg = extract_power(source, "wkg")

        for k in ["15s", "1m", "5m", "20m"]:
            data[f"{k}_w"] = watts.get(k)
            data[f"{k}_wkg"] = wkg.get(k)

        return data

    except Exception as e:
        logging.error(f"❌ Error extracting data: {e}")
        return data

def extract_power(source, category: str):
    pattern = (
        r"<b>\s*(15\s*seconds|1\s*minute|5\s*minutes|20\s*minutes)\s*</b>:\s*"
        r"([\d\.]+)\s*<rsmall>" + re.escape(category)
    )
    matches = re.findall(pattern, source, re.IGNORECASE)
    power_dict = {}
    for label, value in matches:
        label = label.lower()
        if "15" in label:
            power_dict["15s"] = value
        elif "1" in label:
            power_dict["1m"] = value
        elif "5" in label:
            power_dict["5m"] = value
        elif "20" in label:
            power_dict["20m"] = value
    return power_dict

def get_historical_data(filepath):
    column_data = defaultdict(list)

    with open(filepath, mode='r', newline='') as file:
        reader = csv.DictReader(file)
        for row in reader:
            for key, value in row.items():
                column_data[key].append(value)
    
    # Convert defaultdict to regular dict (optional)
    column_data = dict(column_data)
    return column_data

def create_report(history, filename):
    if os.path.exists(filename):
        os.remove(filename)
        logging.info(f"Old version of {filename} exists - deleted")
    xlsx_rp = xlsx_report()
    logging.info(len(history.items()))
    for key, value in history.items():
        xlsx_rp.add_data(key, value)
    xlsx_rp.add_chart("Watts", [3, 6, 7, 8, 9])
    xlsx_rp.add_chart("WKG", [4, 10, 11, 12, 13])
    xlsx_rp.save_file(filename)
    xlsx_rp.add_chart("Racing Score", [5])
    xlsx_rp.save_file(filename)
    logging.info(f'Report created at {filename}')

def build_history(email, password, url, filepath):
    driver = setup_driver()
    try:
        login_to_zwiftpower(driver, email, password)
        data = scrape_profile_data(driver, url)
        pprint(data, sort_dicts=False)
        f_write(filepath, data)
        history = get_historical_data(filepath)
    finally:
        driver.quit()
        logging.info("🛑 Browser closed.")
    return history

def plot_graph(filepath):
    df = pd.read_csv(filepath, parse_dates=['date'])
    os.makedirs('graphs', exist_ok=True)
    plt.figure()
    plt.plot(df['date'], df['zftp'], marker='o')
    plt.title('zFTP')
    plt.xlabel('Date')
    plt.ylabel('zFTP')
    plt.grid(True)
    plt.tight_layout()
    plt.savefig('graphs/zftp_wkg.png')
    plt.close()

def main():
    parser = argparse.ArgumentParser(description="ZwiftPower Scraper with CSV logging and formatted output")
    parser.add_argument("--email", required=True, help="Zwift account email")
    parser.add_argument("--password", required=True, help="Zwift account password")
    parser.add_argument("--url", required=True, help="ZwiftPower profile URL")
    parser.add_argument("--folder", required=True, help="Path to output folder")
    parser.add_argument("--filename", required=True, help="Name for the output file")

    args = parser.parse_args()

    filepath = args.folder + "/" + args.filename
    history = build_history(args.email, args.password, args.url, filepath + ".csv")
    plot_graph(filepath + ".csv")
    create_report(history, filepath + ".xlsx")
    logging.info("📊 Scraped Profile Data:")
    


if __name__ == "__main__":
    main()
